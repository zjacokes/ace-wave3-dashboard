"""Aggregate-only, manuscript-style Table 1 support for ACE Wave 3.

The module builds a canonical participant frame from raw REDCap codes, applies
the two approved dashboard cohorts, precomputes study-group/site/sex views, and
exports every format from the same structured result.  Participant identifiers
are used only to validate one-row-per-participant input; they are never embedded
in dashboard or manuscript outputs.
"""

from __future__ import annotations

import csv
import io
import datetime as dt
import hashlib
import html
import json
import math
import os
import re
from collections import OrderedDict

import numpy as np
import pandas as pd

import config as C


SOURCE_RACE_MAP = {
    "1": "American Indian, Native American, Indigenous, First Nation, or Alaska Native",
    "2": "Asian",
    "3": "Native Hawaiian or Other Pacific Islander",
    "4": "Black, African American, Afro-Caribbean, or African",
    "5": "White/Caucasian",
    "6": "More Than One Race",
}

SOURCE_EDUCATION_MAP = {
    "1": "Did not attend high school",
    "2": "Completed some high school but did not graduate",
    "3": "High school degree or GED",
    "4": "Vocational/Technical school",
    "8": "Some college but did not graduate",
    "9": "Completed an Associate's Degree",
    "10": "Completed a Bachelor's degree",
    "11": ("Completed an advanced degree (Master's Degree and/or a beyond a "
           "Master's degree (e.g., JD, PhD, MD))."),
    "13": "I don't know",
}

SOURCE_GRADE_MAP = {
    "1": "5th grade", "2": "6th grade", "3": "7th grade",
    "4": "8th grade", "5": "9th grade", "6": "10th grade",
    "7": "11th grade", "8": "12th grade", "9": "College",
    "10": "Vocational/Technical school",
}

SOURCE_PERSONAL_INCOME_MAP = {**C.INCOME_MAP, "23": "I don't know"}


def _clean(value) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    return str(value).strip()


def _series(df: pd.DataFrame, field: str) -> pd.Series:
    if field not in df.columns:
        return pd.Series("", index=df.index, dtype="string")
    return df[field].map(_clean).astype("string")


def _parse_number(value):
    raw = _clean(value)
    if not raw:
        return np.nan
    try:
        number = float(raw)
    except (TypeError, ValueError):
        return np.nan
    return number if math.isfinite(number) else np.nan


def _numeric_series(raw: pd.Series):
    values = raw.map(_parse_number).astype(float)
    invalid = raw.ne("") & values.isna()
    return values, invalid


def _mapped_series(raw: pd.Series, mapping: dict):
    values = raw.map(lambda x: mapping.get(_clean(x), "")).astype("string")
    invalid = raw.ne("") & values.eq("")
    return values, invalid


def _collapse_preferred(primary_raw, fallback_raw, mapping):
    """Decode source-specific values, preferring a valid primary response."""
    values = []
    sources = []
    conflicts = []
    invalid_source = []
    for primary, fallback in zip(primary_raw, fallback_raw):
        p = _clean(primary)
        f = _clean(fallback)
        p_label = mapping.get(p)
        f_label = mapping.get(f)
        p_invalid = bool(p and p_label is None)
        f_invalid = bool(f and f_label is None)
        if p_label and f_label:
            values.append(p_label)
            conflicts.append(p_label != f_label)
            sources.append("both_agree" if p_label == f_label else "self_conflict")
        elif p_label:
            values.append(p_label)
            conflicts.append(False)
            sources.append("self")
        elif f_label:
            values.append(f_label)
            conflicts.append(False)
            sources.append("informant")
        else:
            values.append("")
            conflicts.append(False)
            sources.append("none")
        invalid_source.append(p_invalid or f_invalid)
    return (
        pd.Series(values, index=primary_raw.index, dtype="string"),
        pd.Series(sources, index=primary_raw.index, dtype="string"),
        pd.Series(conflicts, index=primary_raw.index, dtype=bool),
        pd.Series(invalid_source, index=primary_raw.index, dtype=bool),
    )


def build_table1_frame(df_raw: pd.DataFrame):
    """Return a canonical participant frame plus input diagnostics."""
    if "record_number" not in df_raw.columns:
        raise ValueError("Table 1 requires record_number to validate participant uniqueness.")

    participant_id = _series(df_raw, "record_number")
    blank_ids = participant_id.eq("")
    duplicate_ids = participant_id.ne("") & participant_id.duplicated(keep=False)
    if blank_ids.any() or duplicate_ids.any():
        problems = []
        if blank_ids.any():
            problems.append(f"{int(blank_ids.sum())} blank record_number value(s)")
        if duplicate_ids.any():
            n_unique = participant_id[duplicate_ids].nunique()
            problems.append(f"{n_unique} duplicated record_number value(s)")
        raise ValueError(
            "Table 1 requires exactly one row per participant; found "
            + " and ".join(problems) + "."
        )

    frame = pd.DataFrame(index=df_raw.index)
    # Retained for validation and cohort selection only.  Exports below never
    # include participant_id.
    frame["participant_id"] = participant_id
    frame["status_raw"] = _series(df_raw, "status")
    frame["inclusion_raw"] = _series(df_raw, "inclusion")
    frame["parent_raw"] = _series(df_raw, "parent")

    age_months_raw = _series(df_raw, "int_age")
    age_months, age_invalid = _numeric_series(age_months_raw)
    frame["age_years"] = age_months / 12.0
    frame["invalid_age_years"] = age_invalid

    sex_raw = _series(df_raw, "sex")
    frame["sex_category"], frame["invalid_sex_category"] = _mapped_series(
        sex_raw, C.VALUE_MAPS["sex"]
    )
    site_raw = _series(df_raw, "site")
    frame["site_category"], frame["invalid_site_category"] = _mapped_series(
        site_raw, C.VALUE_MAPS["site"]
    )
    group_raw = _series(df_raw, "group")
    frame["group_category"], frame["invalid_group_category"] = _mapped_series(
        group_raw, C.VALUE_MAPS["group"]
    )
    asd_raw = _series(df_raw, "asd_group")
    frame["asd_source_category"], frame["invalid_asd_source_category"] = _mapped_series(
        asd_raw, C.VALUE_MAPS["asd_group"]
    )

    race, race_source, race_conflict, race_invalid = _collapse_preferred(
        _series(df_raw, C.RACE_SELF), _series(df_raw, C.RACE_INFORMANT), C.RACE_MAP
    )
    frame["race_category"] = race
    frame["race_source"] = race_source
    frame["race_conflict"] = race_conflict
    frame["invalid_race_category"] = race_invalid

    ethnicity, ethnicity_source, ethnicity_conflict, ethnicity_invalid = (
        _collapse_preferred(
            _series(df_raw, C.HISPANIC_SELF),
            _series(df_raw, C.HISPANIC_INFORMANT),
            C.YESNO,
        )
    )
    frame["ethnicity_category"] = ethnicity
    frame["ethnicity_source"] = ethnicity_source
    frame["ethnicity_conflict"] = ethnicity_conflict
    frame["invalid_ethnicity_category"] = ethnicity_invalid

    age_dx_raw = _series(df_raw, "age_at_dx")
    frame["age_dx_years"], frame["invalid_age_dx_years"] = _numeric_series(age_dx_raw)
    wasi_raw = _series(df_raw, "wasi_fs2_comp")
    frame["wasi_fsiq2"], frame["invalid_wasi_fsiq2"] = _numeric_series(wasi_raw)

    categorical_sources = [
        ("education_category", "ace_demo_17", C.EDUCATION_MAP),
        ("grade_category", "ace_demo_informant_13a", C.GRADE_MAP),
        ("personal_income_category", "ace_demo_26a", C.PERSONAL_INCOME_MAP),
        ("household_income_category", "ace_demo_informant_15", C.INCOME_MAP),
    ]
    for target, source, mapping in categorical_sources:
        frame[target], frame[f"invalid_{target}"] = _mapped_series(
            _series(df_raw, source), mapping
        )

    # Applicability is explicit.  "unknown" is used when this report lacks an
    # upstream branching item; it is never silently treated as applicable.
    frame["app_all"] = "applicable"
    frame["app_asd"] = np.where(
        frame["group_category"].eq("ASD group"), "applicable",
        np.where(frame["group_category"].eq(""), "unknown", "not_applicable"),
    )
    frame["app_adult"] = np.where(
        age_months.ge(228), "applicable",
        np.where(age_months.notna(), "not_applicable", "unknown"),
    )

    informant_eligible = age_months.lt(228) & frame["parent_raw"].eq("1")
    informant_known_no = age_months.ge(228) | frame["parent_raw"].eq("0")
    frame["app_informant_minor"] = np.where(
        informant_eligible, "applicable",
        np.where(informant_known_no, "not_applicable", "unknown"),
    )

    # Current grade additionally branches on a school-attendance field that is
    # not in the report.  A returned response proves applicability; a blank in
    # the informant-eligible population remains unresolved.
    grade_present = frame["grade_category"].ne("") | frame["invalid_grade_category"]
    frame["app_grade"] = np.where(
        grade_present, "applicable",
        np.where(informant_eligible, "unknown",
                 np.where(informant_known_no, "not_applicable", "unknown")),
    )

    # Personal earnings branches on ace_demo_26(1), also absent from this
    # report.  Returned responses are applicable; adult blanks are unresolved.
    personal_present = (
        frame["personal_income_category"].ne("")
        | frame["invalid_personal_income_category"]
    )
    frame["app_personal_income"] = np.where(
        personal_present, "applicable",
        np.where(age_months.ge(228), "unknown",
                 np.where(age_months.notna(), "not_applicable", "unknown")),
    )

    table_fields = {
        "record_number", "status", "inclusion", "parent", "int_age", "sex",
        "site", "group", "asd_group", C.RACE_SELF, C.RACE_INFORMANT,
        C.HISPANIC_SELF, C.HISPANIC_INFORMANT, "age_at_dx", "wasi_fs2_comp",
        "ace_demo_17", "ace_demo_informant_13a", "ace_demo_26a",
        "ace_demo_informant_15",
    }
    absent_fields = sorted(table_fields - set(df_raw.columns))
    diagnostics = {
        "source_rows": int(len(df_raw)),
        "withdrawn_rows": int(frame["status_raw"].eq("1").sum()),
        "race_conflicts": int(frame["race_conflict"].sum()),
        "ethnicity_conflicts": int(frame["ethnicity_conflict"].sum()),
        "invalid_source_values": {
            col.removeprefix("invalid_"): int(frame[col].sum())
            for col in frame.columns if col.startswith("invalid_") and frame[col].any()
        },
        "absent_fields": absent_fields,
    }
    return frame, diagnostics


def _cohort_frame(frame: pd.DataFrame, cohort_key: str) -> pd.DataFrame:
    nonwithdrawn = ~frame["status_raw"].eq("1")
    if cohort_key == "all_nonwithdrawn":
        mask = nonwithdrawn
    elif cohort_key == "included":
        mask = nonwithdrawn & frame["inclusion_raw"].eq("1")
    else:
        raise KeyError(f"Unknown Table 1 cohort: {cohort_key}")
    return frame.loc[mask].copy()


def _stratum_columns(cohort: pd.DataFrame, stratum_key: str):
    spec = C.TABLE1_STRATA[stratum_key]
    field = spec["field"]
    columns = [{
        "key": "overall", "label": "Overall", "mask": pd.Series(True, index=cohort.index)
    }]
    for i, category in enumerate(spec["categories"]):
        columns.append({
            "key": f"s{i + 1}", "label": category,
            "mask": cohort[field].eq(category),
        })
    unknown = ~cohort[field].isin(spec["categories"])
    if unknown.any():
        # Reserve the literal "Unknown" column for missing study group, as
        # specified for the default view.  Alternate presets still account for
        # every participant, but name their unmatched columns explicitly.
        missing_label = (
            "Unknown" if stratum_key == "group"
            else f"Missing {spec['label'].lower()}"
        )
        columns.append({"key": "unknown", "label": missing_label, "mask": unknown})
    for column in columns:
        column["n"] = int(column["mask"].sum())
    return columns


def _cell_accounting(cohort, mask, variable):
    subset = cohort.loc[mask]
    app = subset[variable["app"]]
    invalid_col = f"invalid_{variable['key']}"
    invalid = subset[invalid_col] if invalid_col in subset else pd.Series(False, index=subset.index)
    values = subset[variable["key"]]
    applicable = app.eq("applicable")
    not_applicable = app.eq("not_applicable")
    app_unknown = app.eq("unknown")
    if variable["kind"] == "continuous":
        valid = applicable & values.notna() & ~invalid
        missing = applicable & values.isna() & ~invalid
    else:
        valid = applicable & values.isin(variable["categories"]) & ~invalid
        missing = applicable & values.eq("") & ~invalid
    return {
        "subset": subset,
        "values": values,
        "applicable": applicable,
        "valid": valid,
        "invalid": applicable & invalid,
        "missing": missing,
        "not_applicable": not_applicable,
        "app_unknown": app_unknown,
    }


def _base_cell(variable, column, accounting):
    return {
        "cohort_n": int(column["n"]),
        "applicable_n": int(accounting["applicable"].sum()),
        "valid_n": int(accounting["valid"].sum()),
        "missing_n": int(accounting["missing"].sum()),
        "invalid_n": int(accounting["invalid"].sum()),
        "not_applicable_n": int(accounting["not_applicable"].sum()),
        "applicability_unknown_n": int(accounting["app_unknown"].sum()),
        "scoring_excluded_n": 0,
        "footnote_ids": "|".join(variable.get("footnotes", [])),
    }


def _continuous_cell(cohort, variable, column):
    acc = _cell_accounting(cohort, column["mask"], variable)
    cell = _base_cell(variable, column, acc)
    values = acc["values"].loc[acc["valid"]].astype(float)
    n = len(values)
    stats = {"mean": None, "sd": None, "median": None, "q1": None,
             "q3": None, "min": None, "max": None}
    if n:
        stats.update({
            "mean": float(values.mean()),
            "sd": float(values.std(ddof=1)) if n > 1 else None,
            "median": float(values.quantile(0.5, interpolation="linear")),
            "q1": float(values.quantile(0.25, interpolation="linear")),
            "q3": float(values.quantile(0.75, interpolation="linear")),
            "min": float(values.min()),
            "max": float(values.max()),
        })
    decimals = variable.get("decimals", 1)
    if not n:
        display = "—; n=0"
    elif stats["sd"] is None:
        display = f"{stats['mean']:.{decimals}f} (—); n={n}"
    else:
        display = f"{stats['mean']:.{decimals}f} ({stats['sd']:.{decimals}f}); n={n}"
    cell.update(stats)
    cell.update({"count": None, "denominator": None, "percent": None,
                 "display_value": display})
    return cell


def _categorical_cells(cohort, variable, column):
    acc = _cell_accounting(cohort, column["mask"], variable)
    base = _base_cell(variable, column, acc)
    percent_base = variable.get("percent_base", "column")
    denominator = {
        "column": base["cohort_n"],
        "applicable": base["applicable_n"],
        "valid": base["valid_n"],
    }[percent_base]

    def category_cell(category):
        count = int((acc["valid"] & acc["values"].eq(category)).sum())
        percent = (100.0 * count / denominator) if denominator else None
        display = "—" if not denominator else f"{count} ({percent:.1f})"
        cell = dict(base)
        cell.update({"count": count, "denominator": denominator,
                     "percent": percent, "mean": None, "sd": None,
                     "median": None, "q1": None, "q3": None,
                     "min": None, "max": None, "display_value": display})
        return cell

    if percent_base == "column":
        header = ""
    elif percent_base == "applicable":
        header = f"Applicable n={base['applicable_n']}"
    else:
        unresolved = base["applicability_unknown_n"]
        header = f"Valid n={base['valid_n']}"
        if unresolved:
            header += f"; branch unresolved n={unresolved}"

    cells = {category: category_cell(category) for category in variable["categories"]}
    if percent_base in {"column", "applicable"}:
        missing_count = base["missing_n"] + base["invalid_n"]
        missing_pct = (100.0 * missing_count / denominator) if denominator else None
        missing_cell = dict(base)
        missing_cell.update({
            "count": missing_count, "denominator": denominator,
            "percent": missing_pct, "mean": None, "sd": None,
            "median": None, "q1": None, "q3": None, "min": None, "max": None,
            "display_value": (
                "—" if not denominator else f"{missing_count} ({missing_pct:.1f})"
            ),
        })
        cells["Missing/unavailable"] = missing_cell
    return header, cells


def _footnote_numbers(variables, stratum_key):
    used = {"summary", "withdrawn"}
    for variable in variables:
        used.update(variable.get("footnotes", []))
    if stratum_key == "sex":
        used.add("sex")
    ordered = [key for key in C.TABLE1_FOOTNOTES if key in used]
    return {key: i + 1 for i, key in enumerate(ordered)}


def compute_table1(frame: pd.DataFrame, cohort_key: str, stratum_key: str):
    """Compute one cohort/column-preset table from the canonical frame."""
    cohort = _cohort_frame(frame, cohort_key)
    columns = _stratum_columns(cohort, stratum_key)
    stratum_field = C.TABLE1_STRATA[stratum_key]["field"]
    variables = [v for v in C.TABLE1_VARIABLES if v["key"] != stratum_field]
    footnote_numbers = _footnote_numbers(variables, stratum_key)
    rows = []
    numeric_rows = []

    for variable in variables:
        markers = [footnote_numbers[f] for f in variable.get("footnotes", [])
                   if f in footnote_numbers]
        if variable["kind"] == "continuous":
            row = {"type": "continuous", "label": variable["label"],
                   "variable_key": variable["key"], "markers": markers,
                   "cells": {}}
            for column in columns:
                cell = _continuous_cell(cohort, variable, column)
                row["cells"][column["key"]] = cell
                numeric_rows.append(_numeric_record(
                    cohort_key, stratum_key, variable, None, column, cell
                ))
            rows.append(row)
            continue

        header = {"type": "category_header", "label": variable["label"],
                  "variable_key": variable["key"], "markers": markers,
                  "cells": {}}
        by_column = {}
        for column in columns:
            header_value, cat_cells = _categorical_cells(cohort, variable, column)
            header["cells"][column["key"]] = {"display_value": header_value}
            by_column[column["key"]] = cat_cells
        rows.append(header)

        display_categories = list(variable["categories"])
        if variable.get("percent_base", "column") in {"column", "applicable"}:
            display_categories.append("Missing/unavailable")
        for category in display_categories:
            row = {"type": "category", "label": category,
                   "variable_key": variable["key"], "category": category,
                   "markers": [], "cells": {}}
            for column in columns:
                cell = by_column[column["key"]][category]
                row["cells"][column["key"]] = cell
                numeric_rows.append(_numeric_record(
                    cohort_key, stratum_key, variable, category, column, cell
                ))
            rows.append(row)

    footnotes = [
        {"id": key, "number": number, "text": C.TABLE1_FOOTNOTES[key]}
        for key, number in footnote_numbers.items()
    ]
    return {
        "key": f"{cohort_key}__{stratum_key}",
        "cohort_key": cohort_key,
        "cohort_label": C.TABLE1_COHORTS[cohort_key],
        "stratum_key": stratum_key,
        "stratum_label": C.TABLE1_STRATA[stratum_key]["label"],
        "n": int(len(cohort)),
        "columns": columns,
        "rows": rows,
        "numeric_rows": numeric_rows,
        "footnotes": footnotes,
    }


def _numeric_record(cohort_key, stratum_key, variable, category, column, cell):
    record = {
        "cohort": cohort_key,
        "column_view": stratum_key,
        "variable_key": variable["key"],
        "variable_label": variable["label"],
        "variable_type": variable["kind"],
        "category": category,
        "stratum": column["label"],
        "stratum_key": column["key"],
    }
    for key in (
        "cohort_n", "applicable_n", "valid_n", "missing_n", "invalid_n",
        "not_applicable_n", "applicability_unknown_n", "scoring_excluded_n",
        "count", "denominator", "percent", "mean", "sd", "median", "q1",
        "q3", "min", "max", "display_value", "footnote_ids",
    ):
        record[key] = cell.get(key)
    return record


def _file_sha256(path):
    if not path or not os.path.exists(path):
        return None
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _choices_map(text):
    result = {}
    for piece in _clean(text).split("|"):
        code, sep, label = piece.strip().partition(",")
        if sep and code.strip():
            result[code.strip()] = label.strip()
    return result


def load_dictionary_from_text(text):
    return pd.read_csv(io.StringIO(text.lstrip("\ufeff")), dtype=str, keep_default_na=False)


def load_dictionary(path):
    with open(path, encoding="utf-8-sig") as handle:
        return load_dictionary_from_text(handle.read())


def validate_dictionary(dictionary_path):
    if not dictionary_path:
        return []
    with open(dictionary_path, encoding="utf-8-sig") as handle:
        return validate_dictionary_from_text(handle.read())


def validate_dictionary_from_text(dictionary_text):
    """Compare approved source crosswalks to a supplied REDCap dictionary."""
    if not dictionary_text:
        return []
    dictionary = load_dictionary_from_text(dictionary_text)
    field_col = "Variable / Field Name"
    choice_col = "Choices, Calculations, OR Slider Labels"
    if field_col not in dictionary or choice_col not in dictionary:
        return ["Dictionary columns could not be recognized; crosswalk check skipped."]

    expected = {
        "sex": {"2": "Female", "1": "Male", "3": "Another sex"},
        "site": {"1": "GMU", "2": "Children's National", "3": "UCLA",
                 "4": "Yale", "5": "UVA"},
        "group": {"1": "ASD group", "2": "Unaffected Sibling", "3": "Control"},
        "asd_group": {"1": "ACE Clinic", "2": "Community Dx",
                      "3": "New Diagnostic", "4": "Returning"},
        "ace_demo_13": SOURCE_RACE_MAP,
        "ace_demo_informant_11": SOURCE_RACE_MAP,
        "ace_demo_17": SOURCE_EDUCATION_MAP,
        "ace_demo_informant_13a": SOURCE_GRADE_MAP,
        "ace_demo_26a": SOURCE_PERSONAL_INCOME_MAP,
        "ace_demo_informant_15": C.INCOME_MAP,
    }
    warnings = []
    for field, expected_map in expected.items():
        rows = dictionary.loc[dictionary[field_col].eq(field)]
        if rows.empty:
            warnings.append(f"Dictionary field {field} is absent.")
            continue
        observed = _choices_map(rows.iloc[0][choice_col])
        if observed != expected_map:
            warnings.append(
                f"Dictionary choices for {field} differ from the approved "
                f"{C.TABLE1_CROSSWALK_VERSION} crosswalk."
            )
    return warnings


def build_bundle(df_raw, snapshot_date, status="Draft", source_path=None,
                 dictionary_path=None):
    """Compatibility wrapper for callers that pass disk paths."""
    dictionary_text = None
    if dictionary_path:
        with open(dictionary_path, encoding="utf-8-sig") as handle:
            dictionary_text = handle.read()
    return build_bundle_from_frame(
        df_raw, snapshot_date, status=status, dictionary_text=dictionary_text,
        generated_at=dt.datetime.now().replace(microsecond=0).isoformat(),
        source_hash=_file_sha256(source_path),
        dictionary_hash=_file_sha256(dictionary_path))


def build_bundle_from_frame(df_raw, snapshot_date, *, status="Draft",
                            dictionary_text=None, generated_at,
                            source_hash=None, dictionary_hash=None):
    frame, diagnostics = build_table1_frame(df_raw)
    tables = OrderedDict()
    for cohort_key in C.TABLE1_COHORTS:
        for stratum_key in C.TABLE1_STRATA:
            table = compute_table1(frame, cohort_key, stratum_key)
            tables[table["key"]] = table

    dictionary_warnings = validate_dictionary_from_text(dictionary_text)
    if source_hash is None:
        # API builds have no source file.  Hash a deterministic serialization
        # so every snapshot still has an auditable source fingerprint.
        source_hash = hashlib.sha256(
            df_raw.to_csv(index=False, lineterminator="\n").encode("utf-8")
        ).hexdigest()
    metadata = {
        "title": "ACE Wave 3 participant characteristics",
        "status": status,
        "snapshot_date": snapshot_date,
        "generated_at": generated_at,
        "source_rule": "One unique record_number per row",
        "cohort_rules": {
            "all_nonwithdrawn": "status != 1 (Withdrawn)",
            "included": "status != 1 and inclusion == 1 (Include)",
        },
        "default_view": f"{C.TABLE1_DEFAULT_COHORT}__{C.TABLE1_DEFAULT_STRATUM}",
        "crosswalk_version": C.TABLE1_CROSSWALK_VERSION,
        "source_sha256": source_hash,
        "dictionary_sha256": dictionary_hash,
        "dictionary_warnings": dictionary_warnings,
        "table_specification": {
            "continuous_display": "mean (SD); n=valid",
            "continuous_numeric_export": "n, mean, SD, median, Q1, Q3, min, max",
            "categorical_display": "n (column %); denominator stated for applicable rows",
            "p_values": "not produced",
            "sex_source": "start-page sex only",
            "age_source": "int_age / 12 (age when pre-visit forms were sent)",
        },
        "diagnostics": diagnostics,
    }
    return {
        "tables": tables,
        "default_key": metadata["default_view"],
        "snapshot_date": snapshot_date,
        "status": status,
        "metadata": metadata,
    }


def _marker_html(markers):
    if not markers:
        return ""
    return "<sup>" + ",".join(str(n) for n in markers) + "</sup>"


def render_table(table):
    out = ["<div class='t1-scroll'><table class='table1-table'><thead><tr>",
           "<th scope='col'>Characteristic</th>"]
    for column in table["columns"]:
        out.append(
            f"<th scope='col'>{html.escape(column['label'])}"
            f"<span>N={column['n']}</span></th>"
        )
    out.append("</tr></thead><tbody>")
    for row in table["rows"]:
        classes = f"t1-{row['type'].replace('_', '-')}"
        label = html.escape(row["label"]) + _marker_html(row.get("markers", []))
        scope = "rowgroup" if row["type"] == "category_header" else "row"
        out.append(f"<tr class='{classes}'><th scope='{scope}'>{label}</th>")
        for column in table["columns"]:
            value = row["cells"][column["key"]].get("display_value", "")
            out.append(f"<td>{html.escape(str(value))}</td>")
        out.append("</tr>")
    out.append("</tbody></table></div><ol class='t1-footnotes'>")
    for footnote in table["footnotes"]:
        out.append(
            f"<li value='{footnote['number']}'>{html.escape(footnote['text'])}</li>"
        )
    out.append("</ol>")
    return "".join(out)


def table1_css():
    return """
    section.table1-panel{background:#fff;color:#17242b;border:1px solid #c9d2d7;
      border-radius:4px;box-shadow:0 9px 28px rgba(0,0,0,.24);padding:26px 28px;
      break-inside:auto;page-break-inside:auto}
    section.table1-panel>h2{font-family:Georgia,'Times New Roman',serif;
      color:#17242b;font-size:23px;margin:0 0 4px}
    .t1-kicker{font-size:11px;text-transform:uppercase;letter-spacing:.11em;
      color:#677880;margin:0 0 5px}
    .t1-meta{display:flex;flex-wrap:wrap;gap:7px 18px;margin:12px 0 14px;
      padding:10px 0;border-top:1px solid #d8dfe3;border-bottom:1px solid #d8dfe3;
      font-size:12px;color:#40535c}
    .t1-meta b{color:#17242b}
    .t1-status{display:inline-block;padding:1px 7px;border:1px solid #a86d16;
      color:#80500b;border-radius:12px;font-weight:700;letter-spacing:.03em}
    .t1-note{font-size:12px;color:#485b64;background:#f3f7f7;
      border-left:3px solid #0a7d6c;padding:8px 11px;margin:0 0 14px}
    .t1-controls{display:flex;align-items:end;justify-content:space-between;
      gap:12px;flex-wrap:wrap;margin:0 0 15px}
    .t1-control-groups{display:flex;gap:10px 20px;flex-wrap:wrap}
    .t1-control-group{display:flex;flex-direction:column;gap:4px}
    .t1-control-group>span{font-size:10px;text-transform:uppercase;
      letter-spacing:.08em;color:#64757d;font-weight:700}
    .t1-buttons{display:flex;gap:4px;flex-wrap:wrap}
    .t1-button{appearance:none;background:#fff;color:#31454e;border:1px solid #b9c5ca;
      border-radius:4px;padding:5px 9px;font:600 11px/1.2 -apple-system,BlinkMacSystemFont,
      'Segoe UI',sans-serif;cursor:pointer}
    .t1-button:hover{border-color:#0a7d6c;color:#0a6f60}
    .t1-button.active{background:#0a7d6c;border-color:#0a7d6c;color:#fff}
    .t1-actions{display:flex;gap:6px;align-items:center;flex-wrap:wrap}
    .t1-actions a{font-size:11px;color:#086d60;text-decoration:none;border-bottom:1px solid #8bbab3}
    .t1-copy-state{font-size:10.5px;color:#5b6d75;min-width:42px}
    .table1-view{display:none}.table1-view.active{display:block}
    .t1-scroll{overflow-x:auto;border:1px solid #cbd4d8}
    table.table1-table{border-collapse:collapse;width:100%;min-width:760px;
      font-family:Arial,Helvetica,sans-serif;font-size:11.5px;color:#17242b}
    table.table1-table th,table.table1-table td{border-bottom:1px solid #dce2e5;
      padding:6px 9px;text-align:right;vertical-align:top}
    table.table1-table thead th{background:#edf2f3;color:#263a43;
      font-size:10.5px;text-transform:none;letter-spacing:0;text-align:right;
      border-bottom:2px solid #84979f;white-space:nowrap}
    table.table1-table thead th:first-child{text-align:left;min-width:270px}
    table.table1-table thead th span{display:block;color:#64757d;font-weight:500;
      margin-top:1px}
    table.table1-table tbody th{text-align:left;font-weight:500}
    table.table1-table tr.t1-category-header th,
    table.table1-table tr.t1-category-header td{background:#f5f7f8;
      font-weight:700;border-top:1px solid #879aa2;color:#263a43}
    table.table1-table tr.t1-category-header td{font-size:10px;color:#5b6d75}
    table.table1-table tr.t1-category th{padding-left:25px;color:#354a53}
    table.table1-table tr.t1-continuous th{font-weight:700;border-top:1px solid #879aa2}
    table.table1-table sup{color:#0a6f60;font-size:8px;margin-left:2px}
    .t1-footnotes{font-family:Arial,Helvetica,sans-serif;font-size:10px;
      color:#4b5d65;padding-left:22px;margin:12px 0 0}
    .t1-footnotes li{padding-left:3px;margin:3px 0}
    @media(max-width:760px){section.table1-panel{padding:19px 15px}}
    @media print{
      .t1-controls{display:none!important}.table1-view{display:none!important}
      .table1-view.active{display:block!important}
      section.table1-panel{box-shadow:none;border-color:#9aa7ad;padding:12px;
        break-inside:auto;page-break-inside:auto}
      table.table1-table{font-size:9px;min-width:0}
      table.table1-table th,table.table1-table td{padding:4px 5px}
      table.table1-table tr{break-inside:avoid;page-break-inside:avoid}
      .t1-footnotes{font-size:8.5px}
    }
    """


def render_panel(bundle, export_links=None, print_mode=False):
    export_links = export_links or {}
    default_key = bundle["default_key"]
    default = bundle["tables"][default_key]
    out = ["<section class='table1-panel' id='table1'>",
           "<p class='t1-kicker'>Manuscript-style descriptive table</p>",
           "<h2>Table 1. ACE Wave 3 participant characteristics</h2>",
           "<div class='t1-meta'>",
           f"<span>Cohort: <b id='t1-cohort-label'>{html.escape(default['cohort_label'])}</b></span>",
           f"<span>N: <b id='t1-n'>{default['n']}</b></span>",
           f"<span>Snapshot date: <b>{html.escape(bundle['snapshot_date'])}</b></span>",
           f"<span>Status: <b class='t1-status'>{html.escape(bundle['status'])}</b></span>",
           "</div>",
           "<p class='t1-note'><b>Withdrawal rule:</b> Participants marked "
           "Withdrawn are excluded from both views. The All view therefore "
           "means all non-withdrawn records.</p>"]

    if not print_mode:
        out.extend(["<div class='t1-controls'><div class='t1-control-groups'>",
                    "<div class='t1-control-group'><span>Cohort</span><div class='t1-buttons'>"])
        for key, label in C.TABLE1_COHORTS.items():
            active = " active" if key == C.TABLE1_DEFAULT_COHORT else ""
            short = "All (non-withdrawn)" if key == "all_nonwithdrawn" else "Included only"
            out.append(
                f"<button class='t1-button t1-cohort{active}' data-value='{key}' "
                f"type='button' title='{html.escape(label)}'>{html.escape(short)}</button>"
            )
        out.append("</div></div><div class='t1-control-group'><span>Columns</span><div class='t1-buttons'>")
        for key, spec in C.TABLE1_STRATA.items():
            active = " active" if key == C.TABLE1_DEFAULT_STRATUM else ""
            out.append(
                f"<button class='t1-button t1-stratum{active}' data-value='{key}' "
                f"type='button'>{html.escape(spec['label'])}</button>"
            )
        out.append("</div></div></div><div class='t1-actions'>")
        out.append("<button class='t1-button' id='t1-copy' type='button'>Copy table</button>"
                   "<span class='t1-copy-state' id='t1-copy-state'></span>")
        link_labels = [("standalone", "Standalone HTML"), ("docx", "Word"),
                       ("xlsx", "Excel"), ("numeric", "Numeric CSV")]
        for key, label in link_labels:
            if export_links.get(key):
                out.append(
                    f"<a href='{html.escape(export_links[key])}'>{html.escape(label)}</a>"
                )
        out.append("</div></div>")

    for key, table in bundle["tables"].items():
        if print_mode and key != default_key:
            continue
        active = " active" if key == default_key else ""
        out.append(
            f"<div class='table1-view{active}' data-key='{key}' "
            f"data-cohort-label='{html.escape(table['cohort_label'])}' "
            f"data-n='{table['n']}'>{render_table(table)}</div>"
        )
    if not print_mode:
        out.append(_panel_script())
    out.append("</section>")
    return "".join(out)


def _panel_script():
    return """
    <script>
    (function(){
      var cohort='all_nonwithdrawn', stratum='group';
      function activate(){
        var key=cohort+'__'+stratum;
        document.querySelectorAll('.table1-view').forEach(function(el){
          el.classList.toggle('active', el.dataset.key===key);
        });
        document.querySelectorAll('.t1-cohort').forEach(function(el){
          el.classList.toggle('active', el.dataset.value===cohort);
        });
        document.querySelectorAll('.t1-stratum').forEach(function(el){
          el.classList.toggle('active', el.dataset.value===stratum);
        });
        var view=document.querySelector('.table1-view[data-key="'+key+'"]');
        if(view){
          document.getElementById('t1-cohort-label').textContent=view.dataset.cohortLabel;
          document.getElementById('t1-n').textContent=view.dataset.n;
        }
      }
      document.querySelectorAll('.t1-cohort').forEach(function(el){
        el.addEventListener('click',function(){cohort=el.dataset.value;activate();});
      });
      document.querySelectorAll('.t1-stratum').forEach(function(el){
        el.addEventListener('click',function(){stratum=el.dataset.value;activate();});
      });
      var copy=document.getElementById('t1-copy');
      if(copy){copy.addEventListener('click',function(){
        var table=document.querySelector('.table1-view.active table');
        if(!table){return;}
        var text=Array.from(table.rows).map(function(row){
          return Array.from(row.cells).map(function(cell){
            return cell.innerText.replace(/\\s+/g,' ').trim();
          }).join('\\t');
        }).join('\\n');
        var state=document.getElementById('t1-copy-state');
        function done(){state.textContent='Copied';setTimeout(function(){state.textContent='';},1800);}
        if(navigator.clipboard && navigator.clipboard.writeText){
          navigator.clipboard.writeText(text).then(done);
        }else{
          var area=document.createElement('textarea');area.value=text;
          document.body.appendChild(area);area.select();document.execCommand('copy');
          area.remove();done();
        }
      });}
    })();
    </script>
    """


def display_frame(table):
    columns = [column["label"] for column in table["columns"]]
    records = []
    for row in table["rows"]:
        label = row["label"]
        if row["type"] == "category":
            label = "  " + label
        records.append({
            "Characteristic": label,
            **{
                column["label"]: row["cells"][column["key"]].get("display_value", "")
                for column in table["columns"]
            },
        })
    return pd.DataFrame(records, columns=["Characteristic"] + columns)


def crosswalk_rows():
    specifications = [
        ("sex", {"1": "Male", "2": "Female", "3": "Another sex"},
         C.VALUE_MAPS["sex"]),
        ("site", C.VALUE_MAPS["site"], C.VALUE_MAPS["site"]),
        ("group", {"1": "ASD group", "2": "Unaffected Sibling", "3": "Control"},
         C.VALUE_MAPS["group"]),
        ("asd_group", C.VALUE_MAPS["asd_group"], C.VALUE_MAPS["asd_group"]),
        ("ace_demo_13", SOURCE_RACE_MAP, C.RACE_MAP),
        ("ace_demo_informant_11", SOURCE_RACE_MAP, C.RACE_MAP),
        ("ace_demo_12", {"1": "Yes", "0": "No"}, C.YESNO),
        ("ace_demo_informant_10", {"1": "Yes", "0": "No"}, C.YESNO),
        ("ace_demo_17", SOURCE_EDUCATION_MAP, C.EDUCATION_MAP),
        ("ace_demo_informant_13a", SOURCE_GRADE_MAP, C.GRADE_MAP),
        ("ace_demo_26a", SOURCE_PERSONAL_INCOME_MAP, C.PERSONAL_INCOME_MAP),
        ("ace_demo_informant_15", C.INCOME_MAP, C.INCOME_MAP),
    ]
    rows = []
    for source_field, source_map, publication_map in specifications:
        for order, (raw_code, source_label) in enumerate(source_map.items(), start=1):
            publication_label = publication_map[raw_code]
            canonical = re.sub(r"[^A-Z0-9]+", "_", publication_label.upper()).strip("_")
            rows.append({
                "project": "ACE", "wave": "3",
                "dictionary_version": "2026-09-21",
                "source_field": source_field, "raw_code": raw_code,
                "source_label": source_label, "canonical_value": canonical,
                "publication_label": publication_label, "display_order": order,
                "missing_semantics": "blank = missing/unavailable",
                "rule_version": C.TABLE1_CROSSWALK_VERSION,
            })
    return rows


def _standalone_html(bundle, export_links):
    return (
        "<!doctype html><html lang='en'><head><meta charset='utf-8'>"
        "<meta name='viewport' content='width=device-width,initial-scale=1'>"
        "<title>ACE Wave 3 Table 1</title><style>"
        "*{box-sizing:border-box}body{margin:0;background:#eef2f3;padding:30px;"
        "font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',sans-serif}"
        ".standalone{max-width:1120px;margin:0 auto}"
        + table1_css()
        + "</style></head><body><main class='standalone'>"
        + render_panel(bundle, export_links=export_links, print_mode=False)
        + "</main></body></html>"
    )


def _write_docx(path, bundle):
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    table = bundle["tables"][bundle["default_key"]]
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = section.right_margin = Inches(0.5)
    section.top_margin = section.bottom_margin = Inches(0.5)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title_style = document.styles["Title"]
    title_style.font.name = "Arial"
    title_style.font.size = Pt(18)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title_ppr = title_style._element.pPr
    if title_ppr is not None:
        title_border = title_ppr.find(qn("w:pBdr"))
        if title_border is not None:
            title_ppr.remove(title_border)

    title = document.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(2)
    title.add_run("ACE Wave 3 Participant Characteristics")
    caption = document.add_paragraph()
    caption.paragraph_format.space_after = Pt(3)
    caption.add_run("Table 1").bold = True
    metadata = document.add_paragraph(
        f"Cohort: {table['cohort_label']}    N={table['n']}    "
        f"Snapshot date: {bundle['snapshot_date']}    Status: {bundle['status']}"
    )
    metadata.paragraph_format.space_after = Pt(7)
    for run in metadata.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = None

    def set_cell(cell, text, width, *, fill=None, bold=False,
                 color="17242B", align=WD_ALIGN_PARAGRAPH.LEFT,
                 font_size=8.5):
        cell.width = width
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_w = tc_pr.find(qn("w:tcW"))
        if tc_w is None:
            tc_w = OxmlElement("w:tcW")
            tc_pr.append(tc_w)
        tc_w.set(qn("w:w"), str(int(width.inches * 1440)))
        tc_w.set(qn("w:type"), "dxa")
        margins = tc_pr.find(qn("w:tcMar"))
        if margins is None:
            margins = OxmlElement("w:tcMar")
            tc_pr.append(margins)
        for side in ("top", "left", "bottom", "right"):
            node = OxmlElement(f"w:{side}")
            node.set(qn("w:w"), "70")
            node.set(qn("w:type"), "dxa")
            margins.append(node)
        if fill:
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), fill)
            tc_pr.append(shading)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = align
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.clear()
        run = paragraph.add_run(text)
        run.bold = bold
        run.font.name = "Arial"
        run.font.size = Pt(font_size)
        run.font.color.rgb = None
        run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        # Direct RGB is needed for white-on-dark headers.
        if color == "FFFFFF":
            run.font.color.rgb = RGBColor(255, 255, 255)
        return paragraph

    characteristic_width = Inches(3.4)
    value_width = Inches(1.3)

    # LibreOffice can intermittently place a repeated header outside the top
    # page margin in very long tables.  Build deliberate page-sized chunks so
    # every page starts with an in-flow header and category continuations stay
    # explicit.  This also makes the Word export stable across renderers.
    chunks = []
    chunk = []
    active_header = None
    for row in table["rows"]:
        limit = 20 if not chunks else 23
        if row["type"] == "category_header" and len(chunk) >= limit - 2:
            chunks.append(chunk)
            chunk = []
        elif len(chunk) >= limit:
            chunks.append(chunk)
            chunk = []
            if active_header is not None and row["type"] == "category":
                continuation = dict(active_header)
                continuation["label"] = active_header["label"] + " continued"
                chunk.append(continuation)
        if row["type"] == "category_header":
            active_header = row
        chunk.append(row)
    if chunk:
        chunks.append(chunk)

    def add_table_chunk(rows):
        word_table = document.add_table(rows=1, cols=1 + len(table["columns"]))
        word_table.autofit = False
        tbl_pr = word_table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            tag = OxmlElement(f"w:{edge}")
            tag.set(qn("w:val"), "single")
            tag.set(qn("w:sz"), "4")
            tag.set(qn("w:space"), "0")
            tag.set(qn("w:color"), "D9D9D9")
            borders.append(tag)
        tbl_pr.append(borders)

        header = word_table.rows[0].cells
        set_cell(header[0], "Characteristic", characteristic_width,
                 fill="31535E", bold=True, color="FFFFFF", font_size=9)
        for i, column in enumerate(table["columns"], start=1):
            set_cell(
                header[i], f"{column['label']}\nN={column['n']}", value_width,
                fill="31535E", bold=True, color="FFFFFF",
                align=WD_ALIGN_PARAGRAPH.CENTER, font_size=8.5,
            )
        header_props = word_table.rows[0]._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        header_props.append(repeat)

        category_index = 0
        for row in rows:
            word_row = word_table.add_row()
            cells = word_row.cells
            prefix = "    " if row["type"] == "category" else ""
            marker = "" if not row.get("markers") else " [" + ",".join(
                str(n) for n in row["markers"]
            ) + "]"
            is_header = row["type"] == "category_header"
            is_continuous = row["type"] == "continuous"
            fill = "E9EFF1" if is_header else None
            if row["type"] == "category":
                category_index += 1
                if category_index % 2 == 0:
                    fill = "F7F9FA"
            set_cell(
                cells[0], prefix + row["label"] + marker, characteristic_width,
                fill=fill, bold=is_header or is_continuous,
            )
            for i, column in enumerate(table["columns"], start=1):
                set_cell(
                    cells[i], str(row["cells"][column["key"]].get("display_value", "")),
                    value_width, fill=fill, bold=is_header,
                    align=WD_ALIGN_PARAGRAPH.CENTER,
                )
            row_props = word_row._tr.get_or_add_trPr()
            no_split = OxmlElement("w:cantSplit")
            row_props.append(no_split)
        return word_table

    for index, rows in enumerate(chunks):
        if index:
            document.add_page_break()
        add_table_chunk(rows)

    document.add_page_break()
    for footnote in table["footnotes"]:
        paragraph = document.add_paragraph(
            f"{footnote['number']}. {footnote['text']}"
        )
        paragraph.paragraph_format.space_after = Pt(1)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            run.font.name = "Arial"
            run.font.size = Pt(8)
            run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    document.save(path)


def _write_excel(path, bundle):
    numeric = pd.DataFrame(
        record for table in bundle["tables"].values() for record in table["numeric_rows"]
    )
    metadata_rows = []
    for key, value in bundle["metadata"].items():
        metadata_rows.append({
            "key": key,
            "value": json.dumps(value, sort_keys=True) if isinstance(value, (dict, list)) else value,
        })
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for key, table in bundle["tables"].items():
            cohort, stratum = key.split("__", 1)
            sheet = ("All" if cohort == "all_nonwithdrawn" else "Included") + "_" + stratum
            display_frame(table).to_excel(writer, sheet_name=sheet[:31], index=False)
        numeric.to_excel(writer, sheet_name="Numeric_long", index=False)
        pd.DataFrame(metadata_rows).to_excel(writer, sheet_name="Metadata", index=False)
        workbook = writer.book
        for sheet in workbook.worksheets:
            sheet.freeze_panes = "B2"
            sheet.auto_filter.ref = sheet.dimensions
            sheet.column_dimensions["A"].width = 52
            for column in sheet.iter_cols(min_col=2, max_col=sheet.max_column):
                letter = column[0].column_letter
                sheet.column_dimensions[letter].width = 24


def export_bundle(bundle, outdir, stamp):
    os.makedirs(outdir, exist_ok=True)
    names = export_names(stamp)
    paths = {key: os.path.join(outdir, name) for key, name in names.items()}
    links = {key: name for key, name in names.items()}

    for key, value in (("numeric", numeric_csv_text(bundle)),
                       ("metadata", audit_json_text(bundle)),
                       ("crosswalk", crosswalk_csv_text())):
        with open(paths[key], "w", encoding="utf-8", newline="") as handle:
            handle.write(value)
    _write_excel(paths["xlsx"], bundle)
    _write_docx(paths["docx"], bundle)
    with open(paths["standalone"], "w", encoding="utf-8") as handle:
        handle.write(_standalone_html(bundle, links))
    return {"paths": paths, "links": links}


def export_names(stamp):
    names = {
        "standalone": f"table1_{stamp}.html",
        "docx": f"table1_{stamp}.docx",
        "xlsx": f"table1_{stamp}.xlsx",
        "numeric": f"table1_numeric_{stamp}.csv",
        "metadata": f"table1_metadata_{stamp}.json",
        "crosswalk": f"table1_crosswalk_{C.TABLE1_CROSSWALK_VERSION}.csv",
    }
    return names


def numeric_csv_text(bundle):
    return pd.DataFrame(record for table in bundle["tables"].values()
                        for record in table["numeric_rows"]).to_csv(index=False)


def audit_json_text(bundle):
    return json.dumps(bundle["metadata"], indent=2, sort_keys=True)


def crosswalk_csv_text():
    stream = io.StringIO(newline="")
    rows = crosswalk_rows()
    writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
    writer.writeheader()
    writer.writerows(rows)
    return stream.getvalue()


def load_crosswalk_from_text(text):
    """Read a crosswalk for inspection; approved computation still uses config."""
    return list(csv.DictReader(io.StringIO(text.lstrip("\ufeff"))))


def load_crosswalk(path):
    with open(path, encoding="utf-8-sig", newline="") as handle:
        return load_crosswalk_from_text(handle.read())
